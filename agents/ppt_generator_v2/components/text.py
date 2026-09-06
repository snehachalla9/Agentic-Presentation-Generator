from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

from ..core.theme import Theme


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    font_size: int = Theme.BODY_SIZE,
    bold: bool = False,
    color=None,
    font_name: str = Theme.FONT,
    align=PP_ALIGN.LEFT,
    vertical_anchor=MSO_ANCHOR.TOP,
    margin: float = 0.0,
):
    """
    Add a consistently styled text box to a slide.

    x, y, w, h are specified in inches.
    """

    # Protect renderer from None values
    text = "" if text is None else str(text)

    box = slide.shapes.add_textbox(
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )

    frame = box.text_frame
    frame.clear()

    # Text behaviour
    frame.word_wrap = True
    frame.vertical_anchor = vertical_anchor

    # Internal spacing
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)

    paragraph = frame.paragraphs[0]

    paragraph.text = text
    paragraph.alignment = align
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)

    font = paragraph.font

    font.name = font_name
    font.size = Pt(font_size)
    font.bold = bold
    font.color.rgb = color or Theme.TEXT

    return box


def add_header(
    slide,
    title: str,
    subtitle: str = "",
):
    """
    Add the standard presentation header.

    Returns title and subtitle shapes so they can
    later be inspected by the QA/validation system.
    """

    title_box = add_text(
        slide=slide,
        text=title,
        x=0.7,
        y=0.40,
        w=11.9,
        h=0.60,
        font_size=Theme.TITLE_SIZE,
        bold=True,
        color=Theme.TEXT,
    )

    subtitle_box = None

    if subtitle:
        subtitle_box = add_text(
            slide=slide,
            text=subtitle,
            x=0.7,
            y=1.03,
            w=11.5,
            h=0.40,
            font_size=Theme.SUBTITLE_SIZE,
            color=Theme.MUTED_TEXT,
        )

    return title_box, subtitle_box