from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches

from ..components.text import add_header, add_text
from ..core.theme import Theme


def render_focus(prs, data):
    """
    Strong central message layout.

    Uses:
    - main_message as dominant content
    - callout as supporting label
    - key_takeaway at bottom

    Good for important concepts and transitions.
    """

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = Theme.BACKGROUND

    content = data.get("content", {})

    title = content.get("title", "")
    subtitle = content.get("subtitle", "")
    main_message = content.get("main_message", "")
    callout = content.get("callout", "")
    takeaway = content.get("key_takeaway", "")

    add_header(
        slide=slide,
        title=title,
        subtitle=subtitle,
    )

    # ---------------------------------
    # Large central card
    # ---------------------------------

    panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(1.15),
        Inches(2.05),
        Inches(11.03),
        Inches(3.65),
    )

    panel.fill.solid()
    panel.fill.fore_color.rgb = Theme.CARD
    panel.line.color.rgb = Theme.BORDER

    # Top accent
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(5.55),
        Inches(2.55),
        Inches(2.20),
        Inches(0.07),
    )

    accent.fill.solid()
    accent.fill.fore_color.rgb = Theme.PRIMARY
    accent.line.fill.background()

    # Main idea
    add_text(
        slide=slide,
        text=main_message or takeaway or title,
        x=2.0,
        y=2.95,
        w=9.33,
        h=1.35,
        font_size=26,
        bold=True,
        color=Theme.TEXT,
        align=PP_ALIGN.CENTER,
        vertical_anchor=MSO_ANCHOR.MIDDLE,
    )

    if callout:
        add_text(
            slide=slide,
            text=callout,
            x=2.3,
            y=4.55,
            w=8.73,
            h=0.55,
            font_size=14,
            color=Theme.MUTED_TEXT,
            align=PP_ALIGN.CENTER,
        )

    if takeaway:
        add_text(
            slide=slide,
            text=takeaway,
            x=1.5,
            y=6.30,
            w=10.33,
            h=0.45,
            font_size=12,
            bold=True,
            color=Theme.PRIMARY,
            align=PP_ALIGN.CENTER,
        )

    return slide