from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches

from ..components.text import add_header, add_text
from ..core.theme import Theme


def render_timeline(prs, data):
    """
    Horizontal timeline layout.

    Uses bullets as milestones.
    Best with 3–5 items.
    """

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = Theme.BACKGROUND

    content = data.get("content", {})

    title = content.get("title", "")
    subtitle = content.get("subtitle", "")
    bullets = content.get("bullets", [])
    takeaway = content.get("key_takeaway", "")

    bullets = [
        str(item).strip()
        for item in bullets
        if item is not None and str(item).strip()
    ][:5]

    add_header(
        slide=slide,
        title=title,
        subtitle=subtitle,
    )

    if not bullets:
        return slide

    count = len(bullets)

    left = 1.0
    right = 12.33

    timeline_y = 3.55

    # ---------------------------------
    # Main timeline
    # ---------------------------------

    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left),
        Inches(timeline_y),
        Inches(right - left),
        Inches(0.05),
    )

    line.fill.solid()
    line.fill.fore_color.rgb = Theme.BORDER
    line.line.fill.background()

    spacing = (
        (right - left) / (count - 1)
        if count > 1
        else 0
    )

    # ---------------------------------
    # Milestones
    # ---------------------------------

    for index, bullet in enumerate(bullets):

        if count == 1:
            center_x = 6.666
        else:
            center_x = left + index * spacing

        node = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(center_x - 0.20),
            Inches(timeline_y - 0.18),
            Inches(0.40),
            Inches(0.40),
        )

        node.fill.solid()
        node.fill.fore_color.rgb = Theme.PRIMARY
        node.line.fill.background()

        # Alternate above/below
        above = index % 2 == 0

        if above:
            text_y = 2.0
            number_y = 2.75
        else:
            text_y = 4.20
            number_y = 3.90

        add_text(
            slide=slide,
            text=f"{index + 1:02d}",
            x=center_x - 0.35,
            y=number_y,
            w=0.70,
            h=0.25,
            font_size=9,
            bold=True,
            color=Theme.PRIMARY,
            align=PP_ALIGN.CENTER,
        )

        add_text(
            slide=slide,
            text=bullet,
            x=center_x - 1.0,
            y=text_y,
            w=2.0,
            h=1.0,
            font_size=12,
            bold=True,
            color=Theme.TEXT,
            align=PP_ALIGN.CENTER,
            vertical_anchor=MSO_ANCHOR.MIDDLE,
        )

    if takeaway:
        add_text(
            slide=slide,
            text=takeaway,
            x=1.3,
            y=6.55,
            w=10.7,
            h=0.35,
            font_size=12,
            bold=True,
            color=Theme.PRIMARY,
            align=PP_ALIGN.CENTER,
        )

    return slide