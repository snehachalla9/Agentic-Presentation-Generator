from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches

from ..components.text import add_header, add_text
from ..core.theme import Theme


def render_split(prs, data):
    """
    Clean two-column layout.

    Left:
        main message + takeaway

    Right:
        supporting bullets

    Uses existing content structure only.
    """

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    # ---------------------------------
    # Background
    # ---------------------------------

    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = Theme.BACKGROUND

    # ---------------------------------
    # Content
    # ---------------------------------

    content = data.get("content", {})

    title = content.get("title", "")
    subtitle = content.get("subtitle", "")
    main_message = content.get("main_message", "")
    takeaway = content.get("key_takeaway", "")
    bullets = content.get("bullets", [])

    bullets = [
        str(item).strip()
        for item in bullets
        if item is not None and str(item).strip()
    ][:4]

    # ---------------------------------
    # Header
    # ---------------------------------

    add_header(
        slide=slide,
        title=title,
        subtitle=subtitle,
    )

    # ---------------------------------
    # Left panel
    # ---------------------------------

    left_panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.7),
        Inches(1.8),
        Inches(4.4),
        Inches(4.75),
    )

    left_panel.fill.solid()
    left_panel.fill.fore_color.rgb = Theme.CARD
    left_panel.line.color.rgb = Theme.BORDER

    # Accent
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(1.0),
        Inches(2.2),
        Inches(0.75),
        Inches(0.06),
    )

    accent.fill.solid()
    accent.fill.fore_color.rgb = Theme.PRIMARY
    accent.line.fill.background()

    add_text(
        slide=slide,
        text=main_message or title,
        x=1.0,
        y=2.55,
        w=3.8,
        h=1.65,
        font_size=23,
        bold=True,
        color=Theme.TEXT,
        vertical_anchor=MSO_ANCHOR.MIDDLE,
    )

    if takeaway:
        add_text(
            slide=slide,
            text="KEY TAKEAWAY",
            x=1.0,
            y=4.55,
            w=1.8,
            h=0.25,
            font_size=9,
            bold=True,
            color=Theme.PRIMARY,
        )

        add_text(
            slide=slide,
            text=takeaway,
            x=1.0,
            y=4.95,
            w=3.7,
            h=1.0,
            font_size=14,
            color=Theme.MUTED_TEXT,
        )

    # ---------------------------------
    # Right bullets
    # ---------------------------------

    right_x = 5.55
    top = 1.9

    count = max(len(bullets), 1)

    available_height = 4.55
    gap = 0.18

    item_height = (
        available_height - gap * (count - 1)
    ) / count

    for index, bullet in enumerate(bullets):

        y = top + index * (
            item_height + gap
        )

        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(right_x),
            Inches(y),
            Inches(7.05),
            Inches(item_height),
        )

        box.fill.solid()
        box.fill.fore_color.rgb = Theme.CARD
        box.line.color.rgb = Theme.BORDER

        # Number badge
        badge = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(right_x + 0.30),
            Inches(y + item_height / 2 - 0.22),
            Inches(0.44),
            Inches(0.44),
        )

        badge.fill.solid()
        badge.fill.fore_color.rgb = Theme.PRIMARY
        badge.line.fill.background()

        add_text(
            slide=slide,
            text=str(index + 1),
            x=right_x + 0.30,
            y=y + item_height / 2 - 0.22,
            w=0.44,
            h=0.44,
            font_size=9,
            bold=True,
            color=Theme.CARD,
            align=PP_ALIGN.CENTER,
            vertical_anchor=MSO_ANCHOR.MIDDLE,
        )

        add_text(
            slide=slide,
            text=bullet,
            x=right_x + 1.0,
            y=y + 0.18,
            w=5.65,
            h=item_height - 0.36,
            font_size=14,
            color=Theme.TEXT,
            vertical_anchor=MSO_ANCHOR.MIDDLE,
        )

    return slide