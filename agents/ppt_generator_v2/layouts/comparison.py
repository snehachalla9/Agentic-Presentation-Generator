from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches

from ..components.text import add_header, add_text
from ..core.theme import Theme


def render_comparison(prs, data):
    """
    Render a two-sided comparison slide.

    Supports:
    1. Explicit comparison data
    2. Existing bullet-only content as fallback
    """

    # -------------------------------------------------
    # 1. Create slide
    # -------------------------------------------------

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    # -------------------------------------------------
    # 2. Background
    # -------------------------------------------------

    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = Theme.BACKGROUND

    # -------------------------------------------------
    # 3. Extract content
    # -------------------------------------------------

    content = data.get("content", {})

    title = content.get("title", "")
    subtitle = content.get("subtitle", "")
    takeaway = content.get("key_takeaway", "")

    bullets = content.get("bullets", [])

    bullets = [
        str(item).strip()
        for item in bullets
        if item is not None and str(item).strip()
    ]

    # -------------------------------------------------
    # 4. Extract comparison-specific data
    # -------------------------------------------------

    comparison = data.get("comparison", {})

    left_title = comparison.get(
        "left_title",
        "Perspective A",
    )

    right_title = comparison.get(
        "right_title",
        "Perspective B",
    )

    left_items = comparison.get(
        "left_items",
        [],
    )

    right_items = comparison.get(
        "right_items",
        [],
    )

    # -------------------------------------------------
    # 5. Fallback for current JSON
    # -------------------------------------------------

    if not left_items and not right_items:
        left_items, right_items = _split_bullets(
            bullets
        )

    # -------------------------------------------------
    # 6. Header
    # -------------------------------------------------

    add_header(
        slide=slide,
        title=title,
        subtitle=subtitle,
    )

    # -------------------------------------------------
    # 7. Comparison columns
    # -------------------------------------------------

    _render_comparison_panel(
        slide=slide,
        panel_title=left_title,
        items=left_items,
        x=0.7,
        y=1.85,
        w=5.65,
        h=4.55,
        accent=Theme.PRIMARY,
    )

    _render_comparison_panel(
        slide=slide,
        panel_title=right_title,
        items=right_items,
        x=6.98,
        y=1.85,
        w=5.65,
        h=4.55,
        accent=Theme.SECONDARY,
    )

    # -------------------------------------------------
    # 8. VS marker
    # -------------------------------------------------

    _add_vs_marker(
        slide=slide,
    )

    # -------------------------------------------------
    # 9. Takeaway
    # -------------------------------------------------

    if takeaway:
        add_text(
            slide=slide,
            text=takeaway,
            x=1.2,
            y=6.75,
            w=10.9,
            h=0.35,
            font_size=12,
            bold=True,
            color=Theme.PRIMARY,
            align=PP_ALIGN.CENTER,
        )

    return slide


# =====================================================
# Comparison panel
# =====================================================

def _render_comparison_panel(
    slide,
    panel_title,
    items,
    x,
    y,
    w,
    h,
    accent,
):
    """
    Render one side of the comparison.
    """

    # -------------------------------------------------
    # Panel background
    # -------------------------------------------------

    panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )

    panel.fill.solid()
    panel.fill.fore_color.rgb = Theme.CARD

    panel.line.color.rgb = Theme.BORDER

    # -------------------------------------------------
    # Accent bar
    # -------------------------------------------------

    accent_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(0.08),
        Inches(h),
    )

    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = accent
    accent_bar.line.fill.background()

    # -------------------------------------------------
    # Panel title
    # -------------------------------------------------

    add_text(
        slide=slide,
        text=panel_title,
        x=x + 0.35,
        y=y + 0.30,
        w=w - 0.70,
        h=0.50,
        font_size=19,
        bold=True,
        color=accent,
    )

    # -------------------------------------------------
    # Items
    # -------------------------------------------------

    clean_items = [
        str(item).strip()
        for item in items
        if item is not None and str(item).strip()
    ]

    clean_items = clean_items[:3]

    if not clean_items:
        add_text(
            slide=slide,
            text="No comparison points available.",
            x=x + 0.35,
            y=y + 1.25,
            w=w - 0.70,
            h=0.80,
            font_size=15,
            color=Theme.MUTED_TEXT,
        )

        return

    start_y = y + 1.15
    item_height = 0.95
    gap = 0.18

    for index, item in enumerate(clean_items):

        item_y = start_y + index * (
            item_height + gap
        )

        # Number
        add_text(
            slide=slide,
            text=f"{index + 1:02d}",
            x=x + 0.35,
            y=item_y,
            w=0.45,
            h=0.30,
            font_size=10,
            bold=True,
            color=accent,
        )

        # Content
        add_text(
            slide=slide,
            text=item,
            x=x + 0.95,
            y=item_y - 0.03,
            w=w - 1.35,
            h=item_height,
            font_size=15,
            color=Theme.TEXT,
            vertical_anchor=MSO_ANCHOR.TOP,
        )


# =====================================================
# VS marker
# =====================================================

def _add_vs_marker(slide):
    """
    Add centre VS badge.
    """

    badge = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(6.34),
        Inches(3.60),
        Inches(0.65),
        Inches(0.65),
    )

    badge.fill.solid()
    badge.fill.fore_color.rgb = Theme.TEXT

    badge.line.fill.background()

    add_text(
        slide=slide,
        text="VS",
        x=6.34,
        y=3.60,
        w=0.65,
        h=0.65,
        font_size=10,
        bold=True,
        color=Theme.CARD,
        align=PP_ALIGN.CENTER,
        vertical_anchor=MSO_ANCHOR.MIDDLE,
    )


# =====================================================
# Bullet fallback
# =====================================================

def _split_bullets(bullets):
    """
    Temporary fallback for current JSON.

    Splits existing bullets between two columns.

    Later the Layout Director will generate true
    semantic comparison pairs.
    """

    if not bullets:
        return [], []

    midpoint = (len(bullets) + 1) // 2

    left = bullets[:midpoint]
    right = bullets[midpoint:]

    return left, right