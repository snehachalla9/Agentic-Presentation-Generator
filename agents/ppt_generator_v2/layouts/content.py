# layouts/content.py
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN


def render_content(prs, data: dict):
    """
    Render a standard content slide with title and bullets.
    """
    slide_layout = prs.slide_layouts[6]  # Blank slide
    slide = prs.slides.add_slide(slide_layout)

    content = data.get("content", {})
    title_text = content.get("title", "Untitled")
    subtitle_text = content.get("subtitle", "")
    bullets = content.get("bullets", [])
    key_takeaway = content.get("key_takeaway", "")

    # ============================================================
    # TITLE
    # ============================================================
    title_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(0.3), Inches(12.333), Inches(0.8)
    )
    title_frame = title_box.text_frame
    title_frame.text = title_text
    title_frame.paragraphs[0].font.size = Pt(32)
    title_frame.paragraphs[0].font.bold = True

    # ============================================================
    # SUBTITLE
    # ============================================================
    if subtitle_text:
        subtitle_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(1.1), Inches(12.333), Inches(0.5)
        )
        subtitle_frame = subtitle_box.text_frame
        subtitle_frame.text = subtitle_text
        subtitle_frame.paragraphs[0].font.size = Pt(18)
        subtitle_frame.paragraphs[0].font.italic = True

    # ============================================================
    # BULLETS
    # ============================================================
    start_y = 1.8 if subtitle_text else 1.4
    
    if bullets:
        bullet_box = slide.shapes.add_textbox(
            Inches(0.8), Inches(start_y), Inches(11.733), Inches(4.5)
        )
        bullet_frame = bullet_box.text_frame
        bullet_frame.word_wrap = True

        for i, bullet in enumerate(bullets[:6]):
            if i == 0 and bullet_frame.paragraphs:
                p = bullet_frame.paragraphs[0]
            else:
                p = bullet_frame.add_paragraph()
            
            # Clean bullet text
            bullet_text = str(bullet).strip()
            for char in ["•", "-", "*", "–", "—"]:
                if bullet_text.startswith(char):
                    bullet_text = bullet_text[1:].strip()
            
            p.text = f"• {bullet_text}"
            p.font.size = Pt(18)
            p.level = 0
            p.space_after = Pt(8)

    # ============================================================
    # KEY TAKEAWAY (bottom)
    # ============================================================
    if key_takeaway:
        takeaway_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(6.8), Inches(12.333), Inches(0.6)
        )
        takeaway_frame = takeaway_box.text_frame
        takeaway_frame.text = f"🎯 Key Takeaway: {key_takeaway}"
        takeaway_frame.paragraphs[0].font.size = Pt(16)
        takeaway_frame.paragraphs[0].font.bold = True
        takeaway_frame.paragraphs[0].alignment = PP_ALIGN.CENTER