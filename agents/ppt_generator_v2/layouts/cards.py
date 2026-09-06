from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches

from ..components.text import add_header, add_text
from ..core.theme import Theme


def render_cards(prs, data):
    """Render 1–4 ideas as responsive cards."""

    # ---------------------------------
    # 1. Create blank slide
    # ---------------------------------

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    # ---------------------------------
    # 2. Background
    # ---------------------------------

    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = Theme.BACKGROUND

    # ---------------------------------
    # 3. Extract content
    # ---------------------------------

    content = data.get("content", {})

    title = content.get("title", "")
    subtitle = content.get("subtitle", "")
    bullets = content.get("bullets", [])
    takeaway = content.get("key_takeaway", "")

    # Remove None / empty bullets
    bullets = [
        str(item).strip()
        for item in bullets
        if item is not None and str(item).strip()
    ]

    # First version supports max 4 cards
    bullets = bullets[:4]

    # ---------------------------------
    # 4. Header
    # ---------------------------------

    add_header(
        slide=slide,
        title=title,
        subtitle=subtitle,
    )

    # ---------------------------------
    # 5. Select composition
    # ---------------------------------

    count = len(bullets)

    if count == 0:
        _render_empty_state(
            slide=slide,
            takeaway=takeaway,
        )

    elif count <= 3:
        _render_horizontal_cards(
            slide=slide,
            bullets=bullets,
        )

    else:
        _render_grid_cards(
            slide=slide,
            bullets=bullets,
        )

    # ---------------------------------
    # 6. Takeaway
    # ---------------------------------

    if takeaway:
        add_text(
            slide=slide,
            text=takeaway,
            x=0.8,
            y=6.75,
            w=11.7,
            h=0.35,
            font_size=12,
            bold=True,
            color=Theme.PRIMARY,
        )

    return slide


# =========================================================
# Horizontal composition
# =========================================================

def _render_horizontal_cards(
    slide,
    bullets,
):
    count = len(bullets)

    left = 0.7
    top = 1.85

    available_width = 11.93
    gap = 0.28

    card_width = (
        available_width
        - gap * (count - 1)
    ) / count

    card_height = 4.45

    for index, bullet in enumerate(bullets):

        x = left + index * (
            card_width + gap
        )

        _add_card(
            slide=slide,
            text=bullet,
            number=index + 1,
            x=x,
            y=top,
            w=card_width,
            h=card_height,
        )


# =========================================================
# 2 × 2 grid composition
# =========================================================

def _render_grid_cards(
    slide,
    bullets,
):
    left = 0.7
    top = 1.75

    gap_x = 0.30
    gap_y = 0.25

    available_width = 11.93

    card_width = (
        available_width - gap_x
    ) / 2

    card_height = 2.25

    for index, bullet in enumerate(bullets):

        row = index // 2
        col = index % 2

        x = left + col * (
            card_width + gap_x
        )

        y = top + row * (
            card_height + gap_y
        )

        _add_card(
            slide=slide,
            text=bullet,
            number=index + 1,
            x=x,
            y=y,
            w=card_width,
            h=card_height,
        )


# =========================================================
# Reusable card
# =========================================================

def _add_card(
    slide,
    text,
    number,
    x,
    y,
    w,
    h,
):
    # Card background
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )

    card.fill.solid()
    card.fill.fore_color.rgb = Theme.CARD

    card.line.color.rgb = Theme.BORDER

    # Card number
    add_text(
        slide=slide,
        text=f"{number:02d}",
        x=x + 0.30,
        y=y + 0.30,
        w=0.60,
        h=0.30,
        font_size=10,
        bold=True,
        color=Theme.PRIMARY,
    )

    # Card content
    add_text(
        slide=slide,
        text=text,
        x=x + 0.30,
        y=y + 0.90,
        w=w - 0.60,
        h=h - 1.20,
        font_size=16,
        bold=True,
        color=Theme.TEXT,
        vertical_anchor=MSO_ANCHOR.TOP,
    )


# =========================================================
# Empty fallback
# =========================================================

def _render_empty_state(
    slide,
    takeaway,
):
    message = (
        takeaway
        or "No supporting points available."
    )

    add_text(
        slide=slide,
        text=message,
        x=2.0,
        y=3.0,
        w=9.3,
        h=1.0,
        font_size=22,
        bold=True,
        color=Theme.TEXT,
    )