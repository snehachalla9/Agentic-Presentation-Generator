from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches

from ..components.text import add_header, add_text
from ..core.theme import Theme


def render_process(prs, data):
    """
    Horizontal process / workflow layout.

    Existing bullets become sequential stages.
    Supports 2–5 steps.
    """

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = Theme.BACKGROUND

    content = data.get("content", {})

    title = content.get("title", "")
    subtitle = content.get("subtitle", "")
    bullets = content.get("bullets", [])
    takeaway = content.get("key_takeaway", "")

    bullets = [
        str(item).strip()
        for item in bullets
        if item is not None and str(item).strip()
    ][:5]

    add_header(
        slide=slide,
        title=title,
        subtitle=subtitle,
    )

    if not bullets:
        add_text(
            slide=slide,
            text=takeaway or "No process steps available.",
            x=2.0,
            y=3.0,
            w=9.3,
            h=1.0,
            font_size=22,
            bold=True,
            color=Theme.TEXT,
            align=PP_ALIGN.CENTER,
        )

        return slide

    # ---------------------------------
    # Layout dimensions
    # ---------------------------------

    count = len(bullets)

    left = 0.75
    right = 12.58

    available_width = right - left

    gap = 0.30

    box_width = (
        available_width
        - gap * (count - 1)
    ) / count

    box_width = min(box_width, 2.65)

    total_width = (
        box_width * count
        + gap * (count - 1)
    )

    start_x = (
        13.333 - total_width
    ) / 2

    top = 2.25
    box_height = 3.55

    # ---------------------------------
    # Render stages
    # ---------------------------------

    for index, bullet in enumerate(bullets):

        x = start_x + index * (
            box_width + gap
        )

        # Connector behind cards
        if index < count - 1:
            connector = slide.shapes.add_shape(
                MSO_SHAPE.CHEVRON,
                Inches(
                    x + box_width - 0.02
                ),
                Inches(3.60),
                Inches(gap + 0.10),
                Inches(0.45),
            )

            connector.fill.solid()
            connector.fill.fore_color.rgb = Theme.PRIMARY
            connector.line.fill.background()

        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x),
            Inches(top),
            Inches(box_width),
            Inches(box_height),
        )

        card.fill.solid()
        card.fill.fore_color.rgb = Theme.CARD
        card.line.color.rgb = Theme.BORDER

        # Step number
        badge = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(
                x + box_width / 2 - 0.32
            ),
            Inches(top + 0.38),
            Inches(0.64),
            Inches(0.64),
        )

        badge.fill.solid()
        badge.fill.fore_color.rgb = Theme.PRIMARY
        badge.line.fill.background()

        add_text(
            slide=slide,
            text=f"{index + 1:02d}",
            x=x + box_width / 2 - 0.32,
            y=top + 0.38,
            w=0.64,
            h=0.64,
            font_size=10,
            bold=True,
            color=Theme.CARD,
            align=PP_ALIGN.CENTER,
            vertical_anchor=MSO_ANCHOR.MIDDLE,
        )

        add_text(
            slide=slide,
            text=bullet,
            x=x + 0.25,
            y=top + 1.35,
            w=box_width - 0.50,
            h=1.75,
            font_size=13,
            bold=True,
            color=Theme.TEXT,
            align=PP_ALIGN.CENTER,
            vertical_anchor=MSO_ANCHOR.MIDDLE,
        )

    # ---------------------------------
    # Takeaway
    # ---------------------------------

    if takeaway:
        add_text(
            slide=slide,
            text=takeaway,
            x=1.2,
            y=6.45,
            w=10.9,
            h=0.40,
            font_size=12,
            bold=True,
            color=Theme.PRIMARY,
            align=PP_ALIGN.CENTER,
        )

    return slide