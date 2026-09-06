from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches

from ..components.text import add_header, add_text
from ..core.theme import Theme


def render_steps(prs, data):
    """
    Clean vertical numbered list.

    Best for:
    - 3–5 concepts
    - stages
    - principles
    - key ideas
    """

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = Theme.BACKGROUND

    content = data.get("content", {})

    title = content.get("title", "")
    subtitle = content.get("subtitle", "")
    bullets = content.get("bullets", [])
    takeaway = content.get("key_takeaway", "")

    bullets = [
        str(item).strip()
        for item in bullets
        if item is not None and str(item).strip()
    ][:5]

    add_header(
        slide=slide,
        title=title,
        subtitle=subtitle,
    )

    if not bullets:
        return slide

    count = len(bullets)

    top = 1.75
    available_height = 4.75
    gap = 0.12

    item_height = (
        available_height
        - gap * (count - 1)
    ) / count

    for index, bullet in enumerate(bullets):

        y = top + index * (
            item_height + gap
        )

        # Number
        add_text(
            slide=slide,
            text=f"{index + 1:02d}",
            x=0.85,
            y=y,
            w=0.75,
            h=item_height,
            font_size=18,
            bold=True,
            color=Theme.PRIMARY,
            align=PP_ALIGN.CENTER,
            vertical_anchor=MSO_ANCHOR.MIDDLE,
        )

        # Divider
        divider = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(1.85),
            Inches(y + 0.12),
            Inches(0.04),
            Inches(max(item_height - 0.24, 0.1)),
        )

        divider.fill.solid()
        divider.fill.fore_color.rgb = Theme.BORDER
        divider.line.fill.background()

        # Content
        add_text(
            slide=slide,
            text=bullet,
            x=2.25,
            y=y,
            w=9.8,
            h=item_height,
            font_size=16,
            bold=True,
            color=Theme.TEXT,
            vertical_anchor=MSO_ANCHOR.MIDDLE,
        )

    if takeaway:
        add_text(
            slide=slide,
            text=takeaway,
            x=2.25,
            y=6.70,
            w=9.5,
            h=0.30,
            font_size=11,
            bold=True,
            color=Theme.PRIMARY,
        )

    return slide