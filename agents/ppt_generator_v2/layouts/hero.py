from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches

from ..components.text import add_header, add_text
from ..core.theme import Theme


def render_hero(prs, data):
    """
    Hero layout.

    Best for slides with:
    - one dominant idea
    - a strong main message
    - one supporting visual
    - one takeaway/callout
    """

    # ---------------------------------
    # 1. Create blank slide
    # ---------------------------------

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    # ---------------------------------
    # 2. Background
    # ---------------------------------

    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = Theme.BACKGROUND

    # ---------------------------------
    # 3. Extract content safely
    # ---------------------------------

    content = data.get("content", {})

    title = content.get("title", "")
    subtitle = content.get("subtitle", "")
    main_message = content.get("main_message", "")
    takeaway = content.get("key_takeaway", "")

    # ---------------------------------
    # 4. Header
    # ---------------------------------

    add_header(
        slide,
        title=title,
        subtitle=subtitle,
    )

    # ---------------------------------
    # 5. Main message
    # ---------------------------------

    add_text(
        slide=slide,
        text=main_message,
        x=0.7,
        y=1.8,
        w=5.4,
        h=1.6,
        font_size=24,
        bold=True,
        color=Theme.TEXT,
        vertical_anchor=MSO_ANCHOR.MIDDLE,
    )

    # ---------------------------------
    # 6. Accent line
    # ---------------------------------

    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.7),
        Inches(3.55),
        Inches(0.8),
        Inches(0.06),
    )

    accent.fill.solid()
    accent.fill.fore_color.rgb = Theme.PRIMARY

    accent.line.fill.background()

    # ---------------------------------
    # 7. Key takeaway
    # ---------------------------------

    if takeaway:

        add_text(
            slide=slide,
            text="KEY TAKEAWAY",
            x=0.7,
            y=3.85,
            w=2.0,
            h=0.3,
            font_size=10,
            bold=True,
            color=Theme.PRIMARY,
        )

        add_text(
            slide=slide,
            text=takeaway,
            x=0.7,
            y=4.20,
            w=5.2,
            h=1.2,
            font_size=17,
            color=Theme.MUTED_TEXT,
        )

    # ---------------------------------
    # 8. Visual area
    # ---------------------------------

    visual = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(6.8),
        Inches(1.75),
        Inches(5.8),
        Inches(4.8),
    )

    visual.fill.solid()
    visual.fill.fore_color.rgb = Theme.CARD

    visual.line.color.rgb = Theme.BORDER

        # ---------------------------------
    # 9. Supporting points in the visual panel
    # ---------------------------------

    bullets = [
        str(b).strip()
        for b in content.get("bullets", [])
        if b and str(b).strip()
    ][:5]

    if bullets:
        add_text(
            slide=slide,
            text="\n\n".join(f"•  {b}" for b in bullets),
            x=7.2,
            y=2.05,
            w=5.0,
            h=4.3,
            font_size=14,
            color=Theme.TEXT,
            vertical_anchor=MSO_ANCHOR.MIDDLE,
        )

    return slide